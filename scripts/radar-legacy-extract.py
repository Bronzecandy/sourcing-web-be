"""
Chuyển file recap làm tay "Game Radar Recap.xlsx" (mỗi kỳ 2–3 sheet) thành JSON để nhập vào lịch sử Radar.

    python scripts/radar-legacy-extract.py "<đường dẫn>/Game Radar Recap.xlsx"

Ghi ra data/radar/legacy/legacy-recaps.json; sau đó chạy `npm run radar:import-legacy` để ghi vào DB.
Dùng openpyxl vì exceljs không đọc được file này (có drawing).
"""
import json
import os
import re
import sys
from datetime import date, datetime

import openpyxl

# Mã kỳ trong tên sheet → ngày. Định dạng mã đổi theo thời gian (tháng 6 viết MDD, từ tháng 7 viết D+M),
# nên ghi rõ từng kỳ thay vì đoán.
PERIODS = {
    "610": "2026-06-10",
    "612": "2026-06-12",
    "615": "2026-06-15",
    "622": "2026-06-22",
    "623": "2026-06-23",
    "629": "2026-06-29",
    "17": "2026-07-01",
    "27": "2026-07-02",
    "37": "2026-07-03",
    "278": "2026-08-27",
    "49": "2026-09-04",
    "119": "2026-09-11",
    "189": "2026-09-18",
    "299": "2026-09-29",
    "110": "2026-10-01",
}

HEADER_ALIASES = {
    "Recommend Next Act": "Recommend Next Action",
}

TIERS = {"high": "High", "mid": "Mid", "medium": "Medium", "low": "Low"}


def board_of(title: str) -> str:
    t = title.lower()
    if "new launch" in t:
        return "new_launch"
    if "growth" in t:
        return "reserve_growth"
    return "reserve_scale"


def code_of(title: str, previous: str | None) -> str | None:
    m = re.search(r"\((\d+)\)?\s*$", title.strip())
    if not m:
        return None
    code = m.group(1)
    if code in PERIODS:
        return code
    # Tên sheet bị cắt ở 31 ký tự, vd "Top Potential 10 New Launch (61" → cùng kỳ với sheet trước.
    if previous and previous.startswith(code):
        return previous
    return None


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def number(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def normalize_action(raw: str) -> str:
    t = raw.strip().lower()
    if not t:
        return "Chưa đánh giá"
    if t.startswith("drop") or t.startswith("bỏ"):
        return "Drop"
    m = re.search(r"test viable\s*[-–—]?\s*(high|medium|mid|low)", t)
    if m:
        return f"Test Viable - {TIERS[m.group(1)]} Potential"
    if "test viable" in t:
        return "Test Viable"
    m = re.search(r"monitor\s*[-–—]?\s*(high|medium|mid|low)", t)
    if m:
        tier = TIERS[m.group(1)]
        return f"Monitor - {'Mid' if tier == 'Medium' else tier} Potential"
    if "monitor" in t:
        return "Monitor"
    if "temporary ignore" in t or t.startswith("ignore") or "tạm bỏ qua" in t:
        return "Temporary Ignore"
    if t.startswith("hold"):
        return "Hold"
    if "không phải live service" in t or "không phù hợp" in t:
        return "Drop"
    return "Khác"


def header_of(ws) -> tuple[int, list[str]] | None:
    for r in range(1, 6):
        values = [text(c.value) for c in ws[r]]
        if values and values[0] == "Game":
            return r, [HEADER_ALIASES.get(v, v) for v in values]
    return None


def app_id_of(cell) -> int | None:
    link = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else ""
    m = re.search(r"/app/(\d+)", link)
    return int(m.group(1)) if m else None


def main(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)
    periods: dict[str, dict] = {}
    last_header: dict[str, list[str]] = {}
    previous_code = None
    for ws in wb.worksheets:
        code = code_of(ws.title, previous_code)
        if not code:
            print(f"bỏ qua sheet {ws.title!r}: không đọc được mã kỳ")
            continue
        previous_code = code
        board = board_of(ws.title)
        found = header_of(ws)
        if found:
            header_row, header = found
            last_header[board] = header
        elif board in last_header:
            header_row, header = 1, last_header[board]
            print(f"sheet {ws.title!r}: không có dòng tiêu đề, dùng tiêu đề sheet {board} trước")
        else:
            print(f"bỏ qua sheet {ws.title!r}: không có dòng tiêu đề")
            continue

        col = {name: i for i, name in enumerate(header) if name}
        period = periods.setdefault(
            code,
            {"date": PERIODS[code], "periodCode": PERIODS[code][5:7] + PERIODS[code][8:10], "sheets": [], "rows": []},
        )
        period["sheets"].append(ws.title.strip())
        rank = 0
        for row in ws.iter_rows(min_row=header_row + 1):
            game = row[0]
            app_id = app_id_of(game)
            if not text(game.value) or app_id is None:
                continue
            rank += 1

            def get(name: str) -> str:
                i = col.get(name)
                return text(row[i].value) if i is not None and i < len(row) else ""

            def cell(name: str):
                i = col.get(name)
                return row[i] if i is not None and i < len(row) else None

            title = re.sub(r"TapTap$", "", text(game.value)).strip()
            raw_action = get("Recommend Next Action") or get("Next Action")
            screening_cell = cell("AI Screening") or cell("Screening Result")
            screening_text = text(screening_cell.value) if screening_cell is not None else ""
            screening_url = (
                screening_cell.hyperlink.target
                if screening_cell is not None and screening_cell.hyperlink and screening_cell.hyperlink.target
                else None
            )
            if not screening_url:
                m = re.search(r"https?://\S+", screening_text)
                screening_url = m.group(0) if m else None
            score_m = re.match(r"\s*(\d+(?:\.\d+)?)", screening_text)
            score_i = col.get("Điểm")
            period["rows"].append(
                {
                    "board": board,
                    "rank": rank,
                    "appId": app_id,
                    "title": title,
                    "titleEn": get("Tên (EN)"),
                    "score": number(row[score_i].value) if score_i is not None else None,
                    "chartRank": get("Hạng") if board != "new_launch" else get("Tag BXH"),
                    "trend": get("Xu hướng"),
                    "insights": get("Insights"),
                    "releaseLabel": get("Ngày ra mắt") or get("Ra mắt").replace("—", ""),
                    "nextActionRaw": raw_action,
                    "nextAction": normalize_action(raw_action),
                    "toolAdjustments": get("Tool Adjustments"),
                    "screeningText": screening_text,
                    "screeningUrl": screening_url,
                    "analysisScore": float(score_m.group(1)) if score_m and screening_url else None,
                    "metricsText": get("Chỉ số") or get("Download"),
                    "genre": get("Genre"),
                    "publisher": get("Dev"),
                    "reason": get("Lý do Test"),
                }
            )

    docx_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "radar", "legacy")
    out = []
    for code, period in sorted(periods.items(), key=lambda kv: kv[1]["date"]):
        docx = f"Executive Summary - Radar {period['periodCode']}.docx"
        period["docx"] = docx if os.path.exists(os.path.join(docx_dir, docx)) else None
        out.append(period)
        boards = {}
        for r in period["rows"]:
            boards[r["board"]] = boards.get(r["board"], 0) + 1
        print(f"{period['date']} (mã {code}): {boards} docx={period['docx']}")

    os.makedirs(docx_dir, exist_ok=True)
    target = os.path.join(docx_dir, "legacy-recaps.json")
    with open(target, "w", encoding="utf-8") as f:
        json.dump({"source": os.path.basename(path), "periods": out}, f, ensure_ascii=False, indent=1)
    print(f"đã ghi {target} ({len(out)} kỳ)")


if __name__ == "__main__":
    main(sys.argv[1])
